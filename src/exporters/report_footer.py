"""
Report footer generator.

Produces a uniform block of production details — use case, tool, models,
token usage, a validation notice and use-case specific details — in the
output language of the run.

`add_word_footer_appendix` is used by the Word exporter.
`generate_markdown_footer` renders the same details as Markdown; the
orchestrator does not append it automatically.
"""

import os
from datetime import datetime
from typing import Optional

from src.about import TOOL_NAME, VERSION
from src.output_language import t
from src.pipeline.models import HarvestContext, TaskState


# ─── Use-case labels ────────────────────────────────────────────────

# use case → catalog key of its label ("" = web or institution research)
USE_CASE_LABELS = {
    "":                  "footer.use_case.research",
    "literature_check":  "footer.use_case.literature_check",
    "literature_review": "footer.use_case.literature_review",
    "literature_finder": "footer.use_case.literature_finder",
    "peer_review":       "footer.use_case.peer_review",
    "grant_proposal":    "footer.use_case.grant_proposal",
    "explainer":         "footer.use_case.explainer",
    "decision_analysis": "footer.use_case.decision_analysis",
    "research_design":   "footer.use_case.research_design",
}


def _lang(ctx) -> Optional[str]:
    return getattr(ctx, "output_language", None)


def use_case_label(ctx) -> str:
    key = USE_CASE_LABELS.get(ctx.use_case or "")
    return t(key, _lang(ctx)) if key else ctx.use_case


# ─── Helper functions ────────────────────────────────────────────────

def _format_duration(seconds: float, lang: Optional[str] = None) -> str:
    """Format seconds as e.g. '2 min 05 s' in the output language."""
    if seconds < 60:
        return t("footer.duration_seconds", lang, seconds=int(seconds))
    return t("footer.duration_minutes", lang,
             minutes=int(seconds // 60), seconds=f"{int(seconds % 60):02d}")


def _get_tool_url() -> str:
    """Return the public URL of the tool ('' if not configured)."""
    from src.institution import tool_url
    return tool_url()


def _llm_summary(ctx: HarvestContext) -> dict:
    """Aggregate LLM statistics from ctx.llm_usage."""
    usage = getattr(ctx, "llm_usage", {}) or {}
    primary = usage.get("primary", {}) or {}
    harvest = usage.get("harvest", {}) or {}
    unknown = t("footer.unknown", _lang(ctx))
    return {
        "primary_model": primary.get("model")
                         or os.environ.get("LLM_MODEL_NAME", unknown),
        "harvest_model": harvest.get("model")
                         or os.environ.get("HARVEST_LLM_MODEL_NAME", unknown),
        "primary_calls": primary.get("requests", 0),
        "harvest_calls": harvest.get("requests", 0),
        "total_calls": primary.get("requests", 0) + harvest.get("requests", 0),
        "input_tokens": primary.get("prompt_tokens", 0) + harvest.get("prompt_tokens", 0),
        "output_tokens": primary.get("completion_tokens", 0) + harvest.get("completion_tokens", 0),
    }


def _task_stats(ctx: HarvestContext) -> Optional[dict]:
    """Return sub-task statistics (analysis modes only)."""
    if not ctx.task_results:
        return None
    results = ctx.task_results.values()
    return {
        "total": len(ctx.task_results),
        "done": sum(1 for r in results if r.state == TaskState.DONE),
        "failed": sum(1 for r in results if r.state == TaskState.FAILED),
        "skipped": sum(1 for r in results if r.state == TaskState.SKIPPED),
        "task_tokens": sum(r.tokens_used for r in results),
    }


def _lines(text) -> list[str]:
    if isinstance(text, (list, tuple)):
        return [str(x).strip() for x in text if str(x).strip()]
    return [ln.strip(" -*•\t") for ln in str(text or "").splitlines() if ln.strip(" -*•\t")]


# ─── Use-case specific footer blocks ─────────────────────────────

def _use_case_block(ctx: HarvestContext) -> dict:
    """Use-case specific footer content as a {label: value} dict.

    Used by the Markdown and the Word renderer alike. Reads the preflight
    inputs and the task results of the analysis pipeline.
    """
    from src.pipeline.analysis_pipeline import (
        DESIGN_PREFERENCES, REVIEW_FOCUS, choice_label,
    )
    lang = _lang(ctx)
    pre = ctx.preflight_data or {}
    results = ctx.task_results or {}
    block: dict = {}

    def put(key, value):
        if value not in (None, "", []):
            block[t(key, lang)] = str(value)

    uc = ctx.use_case
    if uc == "explainer":
        put("footer.topic", pre.get("topic"))
        put("footer.audience", pre.get("audience"))
        if ctx.task_plan:
            concepts = [tk.prompt_params.get("concept") for tk in ctx.task_plan.tasks
                        if tk.phase == "explanation" and tk.prompt_params.get("concept")]
            put("footer.core_concepts", ", ".join(concepts))
    elif uc == "peer_review":
        put("footer.discipline", pre.get("discipline"))
        put("footer.review_focus", choice_label(REVIEW_FOCUS, pre.get("review_focus", "")))
    elif uc == "decision_analysis":
        put("footer.n_options", len(_lines(pre.get("options"))) or None)
        put("footer.n_criteria", len(_lines(pre.get("criteria"))) or None)
    elif uc == "grant_proposal":
        put("footer.funder", pre.get("funder"))
        put("footer.duration_project", pre.get("duration"))
    elif uc == "research_design":
        put("footer.discipline", pre.get("discipline"))
        put("footer.design_preference",
            choice_label(DESIGN_PREFERENCES, pre.get("design_preference", "")))
    elif uc in ("literature_review", "literature_finder"):
        search = results.get("SEARCH")
        if search and isinstance(search.parsed_output, dict):
            put("footer.queries_run", len(search.parsed_output.get("queries") or []) or None)
            put("footer.hits_found", len(search.parsed_output.get("results") or []))
    return block


def _field_rows(ctx: HarvestContext) -> list[tuple[str, str]]:
    """All production details as (label, value) rows — shared by both renderers."""
    lang = _lang(ctx)
    llm = _llm_summary(ctx)
    rows = [
        (t("footer.use_case", lang), use_case_label(ctx)),
        (t("footer.tool", lang), f"{TOOL_NAME} {VERSION}"),
    ]
    if _get_tool_url():
        rows.append(("URL", _get_tool_url()))
    rows += [
        (t("footer.created", lang), datetime.now().strftime(t("footer.datetime_format", lang))),
        (t("footer.duration", lang), _format_duration(ctx.duration_seconds, lang)),
        (t("footer.pipeline_calls", lang),
         f"{llm['total_calls']} (primary: {llm['primary_calls']}, harvest: {llm['harvest_calls']})"),
    ]
    if llm["input_tokens"] or llm["output_tokens"]:
        rows.append((t("footer.tokens", lang),
                     t("footer.tokens_value", lang,
                       input=f"{llm['input_tokens']:,}", output=f"{llm['output_tokens']:,}")))
    rows.append((t("footer.primary_llm", lang), llm["primary_model"]))
    rows.append((t("footer.harvest_llm", lang), llm["harvest_model"]))
    stats = _task_stats(ctx)
    if stats:
        rows.append((t("footer.subtasks", lang), t("footer.subtasks_value", lang, **stats)))
        rows.append((t("footer.subtask_tokens", lang), f"~{stats['task_tokens']:,}"))
    rows.extend(_use_case_block(ctx).items())
    if getattr(ctx, "plan_confirmed", False):
        rows.append((t("footer.plan", lang), t("footer.plan_confirmed", lang)))
    return rows


# ─── Markdown footer ────────────────────────────────────────────────

def generate_markdown_footer(ctx: HarvestContext) -> str:
    """Render the production details as a Markdown block."""
    lang = _lang(ctx)
    lines = ["", "---", "", f"## {t('footer.title', lang)}", ""]
    lines += [f"**{label}:** {value}" for label, value in _field_rows(ctx)]
    lines += ["", f"**{t('footer.disclaimer', lang)}**", ""]
    return "\n".join(lines)


# ─── Word footer (for python-docx) ───────────────────────────────────

def add_word_footer_appendix(doc, ctx: HarvestContext) -> None:
    """Add the production details as a Word appendix (after appendix D).

    Args:
        doc: docx.Document
        ctx: HarvestContext
    """
    from docx.shared import Pt, RGBColor

    lang = _lang(ctx)
    doc.add_page_break()
    doc.add_heading(t("footer.title", lang), level=1)

    p = doc.add_paragraph()
    run = p.add_run(t("footer.intro", lang, tool=TOOL_NAME))
    run.italic = True
    run.font.size = Pt(9)

    fields = _field_rows(ctx)
    table = doc.add_table(rows=len(fields), cols=2)
    table.style = "Light List"
    for i, (key, value) in enumerate(fields):
        row = table.rows[i]
        row.cells[0].text = key
        row.cells[1].text = str(value)
        for cell in row.cells:
            for para in cell.paragraphs:
                for r in para.runs:
                    r.font.size = Pt(9)
        for r in row.cells[0].paragraphs[0].runs:   # first column bold
            r.font.bold = True

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    run = p.add_run(t("footer.disclaimer", lang))
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(120, 120, 120)
